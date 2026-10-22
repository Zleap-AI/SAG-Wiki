"""Linux integration checks. Run only in a disposable root-owned container.
Mount this repository at /source and its compiled Linux binaries at /binaries.
Pass --docling to install the real pinned converter and exercise PDF/Office files.
"""
import argparse
import http.client
import http.server
import json
import os
from pathlib import Path
import shutil
import signal
import socket
import socketserver
import subprocess
import sys
import threading
import time
import uuid

parser = argparse.ArgumentParser()
parser.add_argument('--docling', action='store_true')
parser.add_argument('--serve', action='store_true')
args = parser.parse_args()
assert os.getuid() == 0
VERSION = json.loads(Path('/source/package.json').read_text())['version']
base = Path('/srv/sag-test')
app = base / 'app'
data = base / 'data'
for directory in [app / 'bin', data, base / 'alice', base / 'bob']:
    directory.mkdir(parents=True, exist_ok=True)
for name in ['sag-desktop-fnos', 'sag-fnos-worker', 'sag']:
    shutil.copy2('/binaries/' + name, app / 'bin' / name)
shutil.copytree('/source/dist', app / 'www', dirs_exist_ok=True)
os.chown(data, 1500, 1500)
data.chmod(0o700)
for uid, name in [(1501, 'alice'), (1502, 'bob')]:
    directory = base / name
    os.chown(directory, uid, uid)
    directory.chmod(0o700)
    private = directory / 'private.txt'
    private.write_text(name + '-private-marker')
    private.chmod(0o600)
    os.chown(private, uid, uid)
# App access simulates fnOS' application read/write grant, while MCP has only user groups.
subprocess.run(['setfacl', '-m', 'u:1500:rwx', str(base / 'alice'), str(base / 'bob')], check=True)
subprocess.run(['setfacl', '-m', 'u:1500:rw', str(base / 'alice/private.txt'), str(base / 'bob/private.txt')], check=True)
# Re-runs retain the downloaded runtime but rebuild the disposable library/index.
for user in ['1501','1502']:
    for database in (data/'users'/user).glob('sag.sqlite*'):
        database.unlink()
if (base/'alice/.sag-converted').exists():
    shutil.rmtree(base/'alice/.sag-converted')
revoked = set()
class Handler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *values):
        pass
    def do_POST(self):
        assert self.headers['Authorization'] == 'Bearer local-test-only'
        payload = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        uid = payload['data'].get('uid',1501)
        root = base / ('alice' if uid == 1501 else 'bob')
        if payload['req'] == 'trim.file.getUserAccessibleFolders':
            result = {'paths': [str(root)] if uid in [1501, 1502] else []}
        elif payload['req'] == 'trim.file.checkUserACL':
            result = [{'path': path, 'readable': path not in revoked and Path(path).is_relative_to(root),
                       'writable': path not in revoked, 'deletable': path not in revoked}
                      for path in payload['data']['path']]
        elif payload['req'] == 'trim.file.convertPath':
            result = [{'path':path,'semanticPath':path.replace(str(root),'我的文件/CLI 测试')} for path in payload['data']['path']]
        else:
            raise AssertionError(payload['req'])
        content = json.dumps({'code': 0, 'msg': '', 'data': result}).encode()
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(content)))
        self.end_headers()
        self.wfile.write(content)
class Gateway(socketserver.ThreadingMixIn, socketserver.UnixStreamServer):
    daemon_threads = True
native_path = str(base / 'native.sock')
if Path(native_path).exists():
    Path(native_path).unlink()
gateway = Gateway(native_path, Handler)
os.chmod(native_path, 0o666)
threading.Thread(target=gateway.serve_forever, daemon=True).start()
env = {**os.environ, 'TRIM_APPDEST': str(app), 'TRIM_PKGVAR': str(data), 'TRIM_USERNAME': 'sagapp',
       'TRIM_API_TOKEN': 'local-test-only', 'SAG_FNOS_API_SOCKET': native_path}
script = '/source/packaging/sag-desktop-fnos/cmd/main'
def lifecycle(command):
    subprocess.run(['bash', script, command], env=env, check=True)
class Connection(http.client.HTTPConnection):
    def connect(self):
        self.sock = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
        self.sock.settimeout(180)
        self.sock.connect(str(app / 'app.sock'))
def rpc(command, values=None, uid=1501, expected=200):
    connection = Connection('localhost')
    connection.request('POST', '/app/sag-desktop-fnos/api/rpc', json.dumps({
        'command': command, 'args': values or {}, 'session': '12345678-1234-4234-8234-123456789abc'
    }), {'Content-Type': 'application/json', 'X-Trim-Userid': str(uid)})
    response = connection.getresponse()
    content = response.read()
    connection.close()
    assert response.status == expected, (command, response.status, content.decode())
    return json.loads(content)
def wait_for(check, timeout=20):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if check():
            return
        time.sleep(0.05)
    raise AssertionError('check timed out')
def running(pid):
    path = Path(f'/proc/{pid}/stat')
    return path.exists() and path.read_text().split()[2] != 'Z'
def launch_request(values, peer=1500):
    # Use a real unprivileged peer, so SO_PEERCRED is checked in the broker.
    def identity():
        os.setgroups([])
        os.setgid(peer)
        os.setuid(peer)
    code = '''import socket,sys,json
s=socket.socket(socket.AF_UNIX);s.connect(sys.argv[1]);s.sendall(sys.argv[2].encode()+b'\\n')
f=s.makefile('rb');header=f.readline();print(header.decode().strip(),flush=True)
print(f.read().decode(),end='',flush=True)
'''
    return subprocess.run(['python3', '-c', code, str(app / 'worker.sock'), json.dumps(values)],
                          preexec_fn=identity, capture_output=True, text=True, timeout=15)
# Reinstallation can retain root-owned private data. Startup must recover it.
retained = data/'users/1501'
retained.mkdir(parents=True,exist_ok=True)
retained.chmod(0o700)
(retained/'cli-access.json').write_text('{"enabled":true}')
os.chown(retained,0,0)
os.chown(retained/'cli-access.json',0,0)
# fnOS may also retain the application-wide token registry as root.
retained_remote = data/'remote-access.json'
retained_remote.write_text('{"1501":"retained-token-hash"}')
retained_remote.chmod(0o600)
os.chown(retained_remote,0,0)
try:
    lifecycle('start')
    assert retained.stat().st_uid == 1500 and (retained/'cli-access.json').stat().st_uid == 1500
    assert retained_remote.stat().st_uid == 1500 and retained_remote.stat().st_mode & 0o777 == 0o600
    assert rpc('fnos_remote_status')['enabled']
    lifecycle('status')
    assert Path(f"/proc/{(data / 'app.pid').read_text().strip()}/status").read_text().split('Uid:')[1].split()[0] == '1500'
    assert Path(f"/proc/{(data / 'worker.pid').read_text().strip()}/status").read_text().split('Uid:')[1].split()[0] == '0'
    code = '''import os,json
from pathlib import Path
result={'uid':os.getuid(),'groups':os.getgroups(),'token':os.getenv('TRIM_API_TOKEN'),'key':os.getenv('MCP_TEST_KEY'),'home':os.getenv('HOME'),'own':Path('/srv/sag-test/alice/private.txt').read_text()}
try: Path('/srv/sag-test/bob/private.txt').read_text();result['other']='LEAK'
except PermissionError: result['other']='denied'
try: Path('/srv/sag-test/data/users/1501/credentials.json').read_text();result['secret_store']='LEAK'
except PermissionError: result['secret_store']='denied'
print(json.dumps(result),flush=True)
'''
    result = launch_request({'uid':1501,'command':'python3','args':['-c',code],'env':{'MCP_TEST_KEY':'explicit-only'}})
    assert result.returncode == 0, result.stderr
    lines = result.stdout.splitlines()
    reply, value = json.loads(lines[0]), json.loads(lines[1])
    assert reply['error'] is None
    assert value['uid'] == 1501 and value['token'] is None and value['key'] == 'explicit-only'
    assert value['own'] == 'alice-private-marker' and value['other'] == 'denied' and value['secret_store'] == 'denied'
    assert '1500' not in value['groups']
    assert Path(reply['stderr_path']).stat().st_mode & 0o777 == 0o600
    for invalid in [0, 1500, 99999]:
        result = launch_request({'uid':invalid,'command':'id','args':[],'env':{}})
        assert json.loads(result.stdout.splitlines()[0])['error']
    # Make socket permissions permissive only to test the independent kernel-peer check.
    os.chmod(app / 'worker.sock', 0o666)
    result = launch_request({'uid':1501,'command':'id','args':[],'env':{}}, peer=1502)
    assert not result.stdout.strip()
    os.chmod(app / 'worker.sock', 0o600)
    # Verify protocol initialization through the actual HTTP application and Rust MCP client.
    server = {'id':'linux-test','name':'Linux stdio','transport':'stdio','url':'','command':'python3',
              'args':['/source/tests/mock_mcp.py'],'require_confirmation':True,'has_credentials':False,'tools':None}
    tools = rpc('test_mcp', {'server':server,'credentials':{'env':{'MCP_TEST_SUFFIX':'-private'},'headers':{}}})
    assert {tool['name'] for tool in tools} == {'echo','reverse'}
    for uid in [1501, 1502]:
        settings = rpc('get_settings', uid=uid)
        settings['mcp_servers'] = [server]
        rpc('save_settings', {'settings':settings,'answerKey':'','jevKey':'','mcpCredentials':{
            server['id']:{'env':{'MCP_TEST_SUFFIX':str(uid)},'headers':{}}}}, uid=uid)
    assert (data / 'users/1501').stat().st_mode & 0o777 == 0o700
    assert (data / 'users/1502').stat().st_mode & 0o777 == 0o700
    for uid in [1501,1502]:
        values = json.loads((data / f'users/{uid}/credentials.json').read_text())
        assert json.loads(values['mcp:linux-test'])['env']['MCP_TEST_SUFFIX'] == str(uid)
    # Exercise an actual tools/call over the same socket transport, including explicit env.
    invocation = r'''import socket,sys,json
s=socket.socket(socket.AF_UNIX);s.connect(sys.argv[1]);f=s.makefile('rwb', buffering=0)
def request(value): f.write((json.dumps(value)+'\n').encode());return json.loads(f.readline())
header=request({'uid':1501,'command':'python3','args':['/source/tests/mock_mcp.py'],'env':{'MCP_TEST_SUFFIX':'-fnos'}})
assert header['error'] is None
request({'jsonrpc':'2.0','id':1,'method':'initialize','params':{'protocolVersion':'2024-11-05','capabilities':{},'clientInfo':{'name':'integration','version':'1'}}})
result=request({'jsonrpc':'2.0','id':2,'method':'tools/call','params':{'name':'echo','arguments':{'text':'native-tool'}}})
assert result['result']['content'][0]['text']=='native-tool-fnos'
print('tool invoked')
'''
    invocation_result = subprocess.run(['runuser','-u','sagapp','--','python3','-c',invocation,str(app/'worker.sock')], capture_output=True,text=True,timeout=15)
    assert invocation_result.returncode == 0, invocation_result.stderr
    for uid in [1501,1502]:
        assert len(rpc('test_mcp',{'server':server,'credentials':None},uid=uid))==2
    # Disconnecting the peer must kill the command and its descendants.
    command = "import subprocess,time; p=subprocess.Popen(['sleep','300']);print(p.pid,flush=True);time.sleep(300)"
    def identity():
        os.setgroups([]);os.setgid(1500);os.setuid(1500)
    child_script = '''import socket,sys,json,time
s=socket.socket(socket.AF_UNIX);s.connect(sys.argv[1]);s.sendall((json.dumps({'uid':1501,'command':'python3','args':['-c',sys.argv[2]],'env':{}})+'\\n').encode());f=s.makefile('rb');print(f.readline().decode().strip(),flush=True);print(f.readline().decode().strip(),flush=True);sys.stdin.readline()
'''
    peer = subprocess.Popen(['python3','-c',child_script,str(app / 'worker.sock'),command], preexec_fn=identity,
                            stdin=subprocess.PIPE,stdout=subprocess.PIPE,text=True)
    parent_pid = json.loads(peer.stdout.readline())['pid']
    descendant = int(peer.stdout.readline())
    peer.communicate('\n',timeout=5)
    wait_for(lambda: not running(parent_pid) and not running(descendant))
    print('PASS Linux lifecycle, non-root HTTP service, stdio initialization, credentials, peer UID, two users, descendant cleanup', flush=True)
    # Run the packaged CLI through both real Unix peer credentials and the web gateway.
    def sag(arguments, uid=1501):
        who = 'alice' if uid == 1501 else 'bob' if uid == 1502 else 'sagapp' if uid == 1500 else 'root'
        result = subprocess.run(['runuser','-u',who,'--','env',
            'SAG_FNOS_CLI_SOCKET='+str(app/'cli.sock'),str(app/'bin/sag'),*arguments],
            capture_output=True,text=True,timeout=40)
        return result
    root = str(base/'alice')
    assert sag(['--help']).returncode == 0
    assert sag(['config','--enable']).returncode == 0
    added = sag(['add',root,'--json'])
    assert added.returncode == 0, added.stderr
    checked = sag(['check',root,'--json'])
    assert checked.returncode == 0, checked.stderr
    assert json.loads(checked.stdout)['indexed_files'] == 1
    assert json.loads(sag(['list','--json']).stdout)[0]['path'] == root
    matched = sag(['-F','alice-private-marker',root,'--json'])
    assert matched.returncode == 0 and json.loads(matched.stdout)['matched'], matched.stderr
    assert sag(['-F','alice-private-marker',root,'-q']).returncode == 0
    assert sag(['-F','missing-marker',root,'-q']).returncode == 1
    assert sag(['-F','alice-private-marker',root,'-q'],1502).returncode == 2
    assert sag(['status'],1500).returncode != 0
    assert sag(['status'],0).returncode != 0
    revoked.add(str(base/'alice/private.txt'))
    assert sag(['-F','alice-private-marker',root,'-q']).returncode == 2
    revoked.clear()
    web = rpc('fnos_cli_run',{'command':'sag -F "alice-private-marker" --json','directory':root})
    assert web['code'] == 0 and json.loads(web['stdout'])['matched'], web
    web = rpc('fnos_cli_run',{'command':'sag --version','directory':root})
    assert web['code'] == 0 and VERSION in web['stdout'], web
    # Remote MCP authenticates independently of the fnOS browser gateway.
    token = rpc('fnos_remote_rotate')['token']
    def remote_mcp(message, access_token=token, extra=None):
        connection = http.client.HTTPConnection('127.0.0.1',15679,timeout=30)
        headers = {'Content-Type':'application/json','Accept':'application/json, text/event-stream',
                   'MCP-Protocol-Version':'2025-06-18'}
        if access_token is not None: headers['Authorization'] = 'Bearer '+access_token
        headers.update(extra or {})
        connection.request('POST','/mcp',json.dumps(message),headers)
        response=connection.getresponse();content=response.read();connection.close()
        return response.status,content
    initialized={'jsonrpc':'2.0','id':1,'method':'initialize','params':{'protocolVersion':'2025-06-18','capabilities':{},'clientInfo':{'name':'Linux remote verification','version':'1'}}}
    assert remote_mcp(initialized,None)[0] == 401
    status,content=remote_mcp(initialized)
    assert status == 200 and json.loads(content)['result']['serverInfo']['name']=='SAG Wiki fnOS',content
    call={'jsonrpc':'2.0','id':2,'method':'tools/call','params':{'name':'read','arguments':{'path':root+'/private.txt'}}}
    status,content=remote_mcp(call)
    assert status==200 and b'alice-private-marker' in content,content
    call['params']['arguments']['path']=str(base/'bob/private.txt')
    status,content=remote_mcp(call,extra={'X-Trim-Userid':'1502'})
    assert json.loads(content)['result']['isError'] and b'bob-private-marker' not in content,content
    assert remote_mcp(initialized,extra={'Origin':'https://untrusted.example'})[0]==403
    assert token not in (data/'remote-access.json').read_text()
    rpc('fnos_remote_revoke')
    assert remote_mcp(initialized)[0]==401
    print('PASS remote MCP TCP, handshake, real file read, token auth, user isolation and revocation',flush=True)
    stopped = sag(['service','stop'])
    assert stopped.returncode == 0, stopped.stderr
    assert '未运行' in sag(['service','status']).stdout
    restarted = sag(['service','restart'])
    assert restarted.returncode == 0, restarted.stderr
    assert json.loads(restarted.stdout)['version'] == VERSION
    assert rpc('list_roots')[0]['path'] == root
    # fnOS can assign a home path without creating it before the first SSH login.
    home = Path('/home/alice')
    shutil.rmtree(home)
    installed = rpc('fnos_cli_run', {'command':'sag install cli','directory':root})
    assert installed['code'] == 0, installed
    assert home.stat().st_uid == 1501 and home.stat().st_gid == 1501
    assert home.stat().st_mode & 0o777 == 0o700
    home.chmod(0o750)
    installed = rpc('fnos_cli_run', {'command':'sag install codex','directory':root})
    assert installed['code'] == 0 and home.stat().st_mode & 0o777 == 0o750, installed
    for target in ['cli','codex','claude-code','opencode']:
        result = sag(['install',target])
        assert result.returncode == 0, result.stderr
        if target == 'cli':
            installed = subprocess.run(['runuser','-u','alice','--','/home/alice/.local/bin/sag','--version'],capture_output=True,text=True,timeout=10)
            assert installed.returncode == 0 and VERSION in installed.stdout, installed.stderr
        result = sag(['uninstall',target])
        assert result.returncode == 0, result.stderr
    removed = sag(['remove',root,'--json'])
    assert removed.returncode == 0 and (base/'alice/private.txt').exists(), removed.stderr
    print('PASS packaged CLI, web CLI, shared index, search, exits, native ACL, users, lifecycle and all installations',flush=True)
    # Built-in document conversion runs in the application process without the Docling component.
    assert not rpc('get_docling_status')['installed']
    import zipfile
    builtin_source = base / 'alice/内置转换.docx'
    with zipfile.ZipFile(builtin_source, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr('word/document.xml', '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Linux 内置转换 builtinlinuxmarker</w:t></w:r></w:p></w:body></w:document>')
    os.chown(builtin_source, 1501, 1501)
    builtin_bytes = builtin_source.read_bytes()
    rpc('add_root', {'path':root})
    builtin_output = base / 'alice/.sag-converted/内置转换.docx.md'
    wait_for(lambda: any(f['path'] == str(builtin_output) for f in rpc('list_files')), 120)
    assert 'builtinlinuxmarker' in builtin_output.read_text()
    assert builtin_source.read_bytes() == builtin_bytes
    assert rpc('list_roots')[0]['status'] == 'ready', rpc('list_roots')
    direct = subprocess.run(['runuser','-u','alice','--','cat',str(builtin_output)], capture_output=True,text=True)
    assert direct.returncode == 0 and 'builtinlinuxmarker' in direct.stdout, direct.stderr
    result = rpc('run_query', {'query':{'id':str(uuid.uuid4()),'kind':'search','question':'builtinlinuxmarker','scope':[],'references':[],'model':'','conversation_id':None,'replace_from':None,'page':1}})
    assert result['results'] and result['results'][0]['path'] == str(builtin_output), result
    assert not rpc('get_docling_status')['installed']
    # A source changed after indexing is skipped by search until the index catches up.
    builtin_source.write_bytes(builtin_bytes + b'\0')
    stale = rpc('run_query', {'query':{'id':str(uuid.uuid4()),'kind':'search','question':'alice-private-marker','scope':[],'references':[],'model':'','conversation_id':None,'replace_from':None,'page':1}})
    assert stale['results'], stale
    builtin_source.write_bytes(builtin_bytes)
    assert sag(['remove',root,'--json']).returncode == 0
    print('PASS built-in Word conversion without Docling, user-readable output, search, originals preserved',flush=True)

    if args.serve:
        assert sag(['add',root,'--json']).returncode == 0
        class Proxy(http.server.BaseHTTPRequestHandler):
            def log_message(self,*args): pass
            def do_GET(self): self.forward()
            def do_POST(self): self.forward()
            def forward(self):
                connection = Connection('localhost')
                body = self.rfile.read(int(self.headers.get('Content-Length',0)))
                headers = dict(self.headers)
                headers['X-Trim-Userid'] = '1501'
                connection.request(self.command,self.path,body,headers)
                response = connection.getresponse()
                self.send_response(response.status)
                for key,value in response.getheaders():
                    if key.lower() not in ('transfer-encoding','connection'): self.send_header(key,value)
                self.end_headers()
                try:
                    while chunk := response.read1(65536):
                        self.wfile.write(chunk);self.wfile.flush()
                except (BrokenPipeError,ConnectionResetError): pass
                finally: connection.close()
        print('CLI UI fixture listening on 1427',flush=True)
        http.server.ThreadingHTTPServer(('0.0.0.0',1427),Proxy).serve_forever()
    if args.docling:
        converted = base / 'alice/.sag-converted'
        converted.mkdir(exist_ok=True)
        os.chown(converted, 1500, 1500)
        subprocess.run(['setfacl','-m','u:bob:rwx,m::---',str(converted)],check=True)
        status = rpc('get_docling_status')
        if not status['installed']:
            rpc('install_docling')
            previous = None
            end = time.monotonic() + 1800
            while time.monotonic() < end:
                status = rpc('get_docling_status')
                progress = (status['stage'], status['message'][-200:])
                if progress != previous:
                    print('Docling:', *progress, flush=True)
                    previous = progress
                if not status['running']:
                    break
                time.sleep(3)
            assert status['installed'] and not status['error'], status
        home = data / 'users/1501/docling'
        converter_python = home / 'tools/docling-slim/bin/python'
        generator = '''from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
root=Path('/srv/sag-test/alice')
doc=Document();doc.add_heading('Word verification');doc.add_paragraph('wordnativeconversionmarker');doc.save(root/'report.docx')
wb=Workbook();wb.active.append(['Spreadsheet verification','excelnativeconversionmarker']);wb.save(root/'report.xlsx')
p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='Presentation verification';slide.placeholders[1].text='pptnativeconversionmarker';p.save(root/'report.pptx')
# A self-contained PDF with a real text stream and cross-reference table.
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
stream=b'BT /F1 16 Tf 72 720 Td (PDF verification pdfnativeconversionmarker) Tj ET'
objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\\nstream\\n'+stream+b'\\nendstream')
pdf=b'%PDF-1.4\\n'; offsets=[0]
for i,obj in enumerate(objects,1): offsets.append(len(pdf));pdf+=str(i).encode()+b' 0 obj\\n'+obj+b'\\nendobj\\n'
xref=len(pdf);pdf+=b'xref\\n0 6\\n0000000000 65535 f \\n'+b''.join(f'{offset:010d} 00000 n \\n'.encode() for offset in offsets[1:]);pdf+=f'trailer\\n<< /Size 6 /Root 1 0 R >>\\nstartxref\\n{xref}\\n%%EOF\\n'.encode();(root/'report.pdf').write_bytes(pdf)
'''
        subprocess.run([str(converter_python),'-c',generator],check=True)
        originals = {path: path.read_bytes() for path in (base/'alice').glob('report.*')}
        for path in originals:
            os.chown(path, 1501, 1501)
        subprocess.run([str(converter_python),'-c',"import webrtcvad, resemblyzer; assert not webrtcvad.Vad().is_speech(bytes(320),16000)"],check=True)
        # Enable real conversion via the same RPCs as the UI.
        settings = rpc('get_settings')
        settings['convert_other_formats'] = True
        settings['conversion_workers'] = 2
        rpc('save_settings', {'settings':settings,'answerKey':'','jevKey':'','mcpCredentials':None})
        rpc('add_root', {'path':str(base/'alice')})
        def indexed():
            files = rpc('list_files')
            return len([f for f in files if '.sag-converted/report.' in f['path']]) == 4
        wait_for(indexed, 600)
        for extension, marker in [('pdf','pdfnativeconversionmarker'),('docx','wordnativeconversionmarker'),('xlsx','excelnativeconversionmarker'),('pptx','pptnativeconversionmarker')]:
            output = base / f'alice/.sag-converted/report.{extension}.md'
            assert marker in output.read_text(), (extension, output.read_text())
            document = rpc('read_document', {'path':str(output)})
            assert marker in json.dumps(document)
            direct = subprocess.run(['runuser','-u','alice','--','cat',str(output)], capture_output=True,text=True)
            assert direct.returncode == 0 and marker in direct.stdout, direct.stderr
            result = rpc('run_query', {'query':{'id':str(uuid.uuid4()),'kind':'search','question':marker,'scope':[],'references':[],'model':'','conversation_id':None,'replace_from':None,'page':1}})
            assert result['results'] and marker in json.dumps(result), result
        assert all(path.read_bytes() == content for path, content in originals.items())
        acl = subprocess.check_output(['getfacl','-cp',str(converted)],text=True)
        assert 'user:bob:---' in acl, acl
        # Cached converted content must become inaccessible immediately after source revocation.
        revoked.add(str(base/'alice/report.pdf'))
        rpc('read_document', {'path':str(base/'alice/.sag-converted/report.pdf.md')}, expected=400)
        assert not rpc('get_docling_status', uid=1502)['installed']
        print('PASS real Docling download/install; PDF, Word, Excel, PowerPoint conversion and indexing; originals preserved; source revocation; per-user component state',flush=True)
finally:
    if sys.exc_info()[0]:
        for log in ['app.log','worker.log']:
            path=data/log
            if path.exists(): print(log,path.read_text()[-5000:],flush=True)
    lifecycle('stop')
    assert not (app / 'worker.sock').exists() and not (app / 'app.sock').exists()
    gateway.shutdown()
